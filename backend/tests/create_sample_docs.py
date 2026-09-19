import os
import docx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

def create_sample_docx(path: str):
    doc = docx.Document()
    doc.add_heading("RESIDENTIAL LEASE AGREEMENT", level=0)
    
    doc.add_heading("1. TERM AND AUTO-RENEWAL", level=1)
    doc.add_paragraph(
        "This Lease Agreement shall commence on January 1, 2026, and shall automatically renew for successive "
        "one-year terms unless tenant provides written notice of cancellation at least 90 days prior to expiration."
    )
    
    doc.add_heading("2. RENT MODIFICATION AND LATE FEES", level=1)
    doc.add_paragraph(
        "Landlord reserves the right to modify rent or fees at any time in its sole discretion without prior notice. "
        "A late fee of 15% plus compounding interest will apply to any payment delayed past 5 days."
    )
    
    doc.add_heading("3. LIABILITY WAIVER AND AS-IS CONDITION", level=1)
    doc.add_paragraph(
        "Tenant agrees to waive all liability against Landlord for any injury, loss, or property damage on the premises. "
        "The property is accepted strictly on an as-is basis."
    )

    doc.add_heading("4. DISPUTE RESOLUTION AND ARBITRATION", level=1)
    doc.add_paragraph(
        "All claims and disputes arising under this lease shall be resolved through binding arbitration. "
        "Tenant waives right to a jury trial and agrees to refrain from any class action waiver."
    )

    doc.save(path)

def create_sample_pdf(path: str, title: str, sections: list):
    c = canvas.Canvas(path, pagesize=letter)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, 750, title)
    
    y = 710
    for header, body in sections:
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, y, header)
        y -= 20
        c.setFont("Helvetica", 10)
        c.drawString(50, y, body)
        y -= 35

    c.save()

if __name__ == "__main__":
    out_dir = os.path.join(os.path.dirname(__file__), "sample_docs")
    os.makedirs(out_dir, exist_ok=True)
    
    create_sample_docx(os.path.join(out_dir, "sample_lease.docx"))
    print("Created sample_lease.docx")

    create_sample_pdf(
        os.path.join(out_dir, "sample_tos.pdf"),
        "TERMS OF SERVICE AGREEMENT",
        [
            ("SECTION 1. UNILATERAL MODIFICATIONS", "Company may modify at any time in its sole discretion without notice all terms."),
            ("SECTION 2. TERMINATION AND DATA SHARING", "Company may terminate for any reason at will. Personal data is shared with third parties for any purpose."),
            ("SECTION 3. REFUNDS AND FEES", "All payments under this contract are non-refundable with no refunds granted under any circumstances.")
        ]
    )
    print("Created sample_tos.pdf")

    create_sample_pdf(
        os.path.join(out_dir, "sample_insurance.pdf"),
        "HEALTH & PROPERTY INSURANCE POLICY",
        [
            ("1. COVERAGE EXCLUSIONS", "This policy does not cover damages caused by water ingress, excluded perils, or pre-existing conditions."),
            ("2. INDEMNIFICATION OBLIGATIONS", "Policyholder agrees to indemnify and hold harmless the insurer and defend at your own expense."),
            ("3. PREMIUM INCREASES", "Insurer may increase rent or premiums at landlord's discretion without any cap or limit.")
        ]
    )
    print("Created sample_insurance.pdf")
